"""Build the CryoStack 10-minute launch deck (editable PPTX + PDF preview).

Pipeline:
  1. SVG background per slide (diagrams, cards, images) -> rendered to PNG via rsvg-convert.
     Headline/kicker/central-statement/footer/URL text is left OUT of the SVG so it can
     live as real, editable text boxes in the PPTX (not flattened into the image).
  2. python-pptx assembles the deck: full-bleed background picture + real text boxes
     (with speaker notes).
  3. A PIL compositor draws the same overlay text on top of the same background PNGs to
     produce a faithful visual QA preview and a PDF (no LibreOffice available in this env).
"""
from __future__ import annotations

import subprocess
import textwrap
from dataclasses import dataclass, field
from pathlib import Path
from xml.sax.saxutils import escape

import qrcode
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parents[1]
ASSETS = OUT / "assets"
RENDERS = OUT / "output" / "slides_svg"
PREVIEWS = OUT / "output" / "previews"
for d in (ASSETS, RENDERS, PREVIEWS):
    d.mkdir(parents=True, exist_ok=True)

W, H = 1920, 1080
PX_TO_EMU = 6350  # exact: 12192000/1920 == 6858000/1080

NAVY = "082B59"
BLUE = "087EEB"
CYAN = "24B8E8"
PALE = "EEF7FC"
BG = "F8FBFD"
INK = "172033"
MUTED = "5E7187"
WHITE = "FFFFFF"
GREEN = "177A55"
BORDER = "D6E5F0"
ACCENTS = [BLUE, CYAN, GREEN, NAVY]

FONT_BOLD = "/usr/share/fonts/liberation-sans/LiberationSans-Bold.ttf"
FONT_REG = "/usr/share/fonts/liberation-sans/LiberationSans-Regular.ttf"

SOURCE_LOGO = ROOT / "icesee_jupyter_book" / "cryostack.png"
SOURCE_ISSM_RESULT = ROOT / "advertising" / "cryostack_launch" / "assets" / "issm-stressbalance-velocity.png"
SOURCE_ICESEE_RESULT = ROOT / "external" / "ICESEE" / "scripts" / "ci" / "figures" / "lorenz96_ci.png"

MEDIA = OUT / "media"
DEMO_PANELS = [
    ("issm_remote", "ISSM — Remote HPC"),
    ("issm_cloud", "ISSM — Cloud"),
    ("icepack_cloud", "Icepack — Cloud"),
]

CRYOSTACK_URL = "https://cryostack.eas.gatech.edu/index.html"
DOCS_URL = "cryostack.eas.gatech.edu/documentation.html"


def hexrgb(h: str) -> RGBColor:
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def prepare_assets() -> None:
    import shutil
    from PIL import ImageChops
    shutil.copy2(SOURCE_ISSM_RESULT, ASSETS / "issm-stressbalance-velocity.png")
    shutil.copy2(SOURCE_ICESEE_RESULT, ASSETS / "icesee-lorenz96-result.png")

    # Isolate the CryoStack hexagon mark (above the wordmark) from the canonical logo,
    # then key out its near-white background so it composites cleanly on any slide bg.
    logo = Image.open(SOURCE_LOGO).convert("RGB")
    region = logo.crop((105, 20, 1435, 600))
    bg = Image.new("RGB", region.size, (255, 255, 255))
    bbox = ImageChops.difference(region, bg).getbbox()
    tight = region.crop(bbox)
    pad = 24
    canvas = Image.new("RGB", (tight.size[0] + 2 * pad, tight.size[1] + 2 * pad), (255, 255, 255))
    canvas.paste(tight, (pad, pad))
    rgba = canvas.convert("RGBA")
    data = rgba.getdata()
    new_data = [
        (r, g, b, 0) if r > 248 and g > 248 and b > 248 else (r, g, b, a)
        for (r, g, b, a) in data
    ]
    rgba.putdata(new_data)
    rgba.save(ASSETS / "cryostack-icon-only.png")

    qr = qrcode.QRCode(border=4, box_size=14)  # 4-module quiet zone per QR spec, for reliable scanning
    qr.add_data(CRYOSTACK_URL)
    qr.make(fit=True)
    img = qr.make_image(fill_color=NAVY and "#" + NAVY, back_color="white").convert("RGB")
    img.save(ASSETS / "qr-cryostack.png")


def demo_thumbnail_path(folder: str) -> Path:
    """Real recording frame if it exists; otherwise the generated placeholder. Never fabricated UI."""
    real = MEDIA / folder / "thumbnail.png"
    if real.exists():
        return real
    return MEDIA / folder / "thumbnail_placeholder.png"


def prepare_media_placeholders() -> None:
    """Create media/<name>/ with a clearly-labeled placeholder thumbnail, unless a real one exists."""
    pw, ph = 1280, 720
    pale = tuple(int(PALE[i:i+2], 16) for i in (0, 2, 4))
    border = tuple(int(BORDER[i:i+2], 16) for i in (0, 2, 4))
    muted = tuple(int(MUTED[i:i+2], 16) for i in (0, 2, 4))
    for folder, label in DEMO_PANELS:
        d = MEDIA / folder
        d.mkdir(parents=True, exist_ok=True)
        if (d / "thumbnail.png").exists():
            continue  # a real recording frame has been dropped in -- never overwrite it
        im = Image.new("RGB", (pw, ph), pale)
        draw = ImageDraw.Draw(im)
        draw.rectangle([4, 4, pw-5, ph-5], outline=border, width=4)
        font_small = ImageFont.truetype(FONT_BOLD, 30)
        font_tiny = ImageFont.truetype(FONT_REG, 22)
        msg1 = "PLACEHOLDER"
        msg2 = f"recording pending -- {label}"
        w1 = font_small.getlength(msg1)
        w2 = font_tiny.getlength(msg2)
        # Text sits near the bottom, clear of the play-button glyph the slide draws over the center.
        draw.text(((pw-w1)/2, ph - 96), msg1, font=font_small, fill=muted)
        draw.text(((pw-w2)/2, ph - 58), msg2, font=font_tiny, fill=muted)
        im.save(d / "thumbnail_placeholder.png")


# ---------------------------------------------------------------- SVG helpers

def svg_lines(lines: list[str], x: int, y: int, css: str, line_height: int, anchor: str = "start") -> str:
    spans = "".join(
        f'<tspan x="{x}" dy="{0 if i == 0 else line_height}">{escape(line)}</tspan>'
        for i, line in enumerate(lines)
    )
    return f'<text x="{x}" y="{y}" text-anchor="{anchor}" class="{css}">{spans}</text>'


def pill(x: int, y: int, w: int, h: int, label: str, fill=PALE, stroke=CYAN, text_cls="pill", r=999, wrap_chars=18) -> str:
    lines = textwrap.wrap(label, width=wrap_chars)
    ty = y + h / 2 - (len(lines) - 1) * 13
    return (
        f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="{r}" fill="#{fill}" stroke="#{stroke}" stroke-width="2"/>'
        + svg_lines(lines, x + w / 2, ty + 8, text_cls, 26, "middle")
    )


def box(x: int, y: int, w: int, h: int, label: str, sub: str = "", fill=WHITE, stroke=BORDER, label_cls="boxlabel", sub_cls="boxsub", r=16, label_wrap=22, sub_wrap=30) -> str:
    out = f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="{r}" fill="#{fill}" stroke="#{stroke}" stroke-width="2"/>'
    label_lh, sub_lh, gap = 32, 23, 12
    lines = textwrap.wrap(label, width=label_wrap)
    sub_lines = textwrap.wrap(sub, width=sub_wrap) if sub else []
    block_h = len(lines) * label_lh + (gap + len(sub_lines) * sub_lh if sub_lines else 0)
    first_baseline = y + h / 2 - block_h / 2 + label_lh * 0.72
    out += svg_lines(lines, x + w / 2, first_baseline, label_cls, label_lh, "middle")
    if sub_lines:
        sub_baseline = first_baseline + (len(lines) - 1) * label_lh + label_lh * 0.5 + gap + sub_lh * 0.6
        out += svg_lines(sub_lines, x + w / 2, sub_baseline, sub_cls, sub_lh, "middle")
    return out


def vline(x: int, y1: int, y2: int, stroke=CYAN, w=3, dash: str | None = None) -> str:
    d = f' stroke-dasharray="{dash}"' if dash else ""
    return f'<line x1="{x}" y1="{y1}" x2="{x}" y2="{y2}" stroke="#{stroke}" stroke-width="{w}"{d}/>'


def hline(x1: int, x2: int, y: int, stroke=CYAN, w=3, dash: str | None = None) -> str:
    d = f' stroke-dasharray="{dash}"' if dash else ""
    return f'<line x1="{x1}" y1="{y}" x2="{x2}" y2="{y}" stroke="#{stroke}" stroke-width="{w}"{d}/>'


def chevron_flow(labels: list[str], x0: int, y: int, total_w: int, h: int = 96, gap: int = 26) -> str:
    n = len(labels)
    bw = (total_w - gap * (n - 1)) / n
    out = ""
    for i, label in enumerate(labels):
        x = x0 + i * (bw + gap)
        out += pill(int(x), y, int(bw), h, label, fill=PALE, stroke=CYAN, text_cls="flowlabel", r=14)
        if i < n - 1:
            ax = x + bw + gap / 2
            out += f'<text x="{ax}" y="{y + h/2 + 10}" text-anchor="middle" class="arrow">&#8594;</text>'
    return out


def app_icon(kind: str, cx: float, cy: float, s: float, color: str) -> str:
    """Small monochrome line-art glyph suggesting each application's science, not its UI."""
    c = f"#{color}"
    if kind == "mountain":
        return (
            f'<path d="M {cx-s/2} {cy+s/3} L {cx-s/8} {cy-s/2} L {cx+s/10} {cy-s/8} '
            f'L {cx+s/3} {cy-s/2.6} L {cx+s/2} {cy+s/3} Z" fill="none" stroke="{c}" '
            f'stroke-width="4" stroke-linejoin="round" stroke-linecap="round"/>'
        )
    if kind == "ensemble":
        heights = [0.35, 0.65, 1.0, 0.55, 0.3]
        n = len(heights)
        bw = s / (n * 1.8)
        gap = bw * 0.8
        total = n * bw + (n - 1) * gap
        x0 = cx - total / 2
        out = ""
        for i, h in enumerate(heights):
            bh = s * 0.55 * h
            x = x0 + i * (bw + gap)
            out += f'<rect x="{x}" y="{cy + s*0.28 - bh}" width="{bw}" height="{bh}" rx="{bw/2}" fill="{c}" opacity="{0.5+0.5*h:.2f}"/>'
        return out
    if kind == "radar":
        out = ""
        for r, op in ((s*0.5, 0.9), (s*0.34, 0.6), (s*0.18, 0.35)):
            out += f'<path d="M {cx-r} {cy+s*0.12} A {r} {r} 0 0 1 {cx+r} {cy+s*0.12}" fill="none" stroke="{c}" stroke-width="3.5" opacity="{op}"/>'
        out += f'<circle cx="{cx}" cy="{cy+s*0.12}" r="4.5" fill="{c}"/>'
        return out
    if kind == "thermometer":
        stem_w = s * 0.16
        stem_x = cx - stem_w / 2
        top_y = cy - s * 0.5
        bulb_y = cy + s * 0.32
        out = f'<rect x="{stem_x}" y="{top_y}" width="{stem_w}" height="{bulb_y-top_y}" rx="{stem_w/2}" fill="none" stroke="{c}" stroke-width="3.5"/>'
        out += f'<circle cx="{cx}" cy="{bulb_y}" r="{s*0.14}" fill="{c}"/>'
        out += f'<rect x="{stem_x+stem_w*0.28}" y="{cy-s*0.05}" width="{stem_w*0.44}" height="{bulb_y-(cy-s*0.05)}" fill="{c}"/>'
        return out
    return ""


def play_button(cx: float, cy: float, r: float, bg=NAVY, fg=WHITE, opacity: float = 0.82) -> str:
    tri = r * 0.42
    points = f"{cx-tri*0.55},{cy-tri} {cx-tri*0.55},{cy+tri} {cx+tri*0.9},{cy}"
    return (
        f'<circle cx="{cx}" cy="{cy}" r="{r}" fill="#{bg}" opacity="{opacity}"/>'
        f'<polygon points="{points}" fill="#{fg}"/>'
    )


def image_tag(path: Path, x: int, y: int, w: int, h: int) -> str:
    import base64
    b64 = base64.b64encode(path.read_bytes()).decode()
    return f'<image href="data:image/png;base64,{b64}" x="{x}" y="{y}" width="{w}" height="{h}" preserveAspectRatio="xMidYMid meet"/>'


STYLE = f'''<style>
text{{font-family:'DejaVu Sans',Arial,sans-serif}}
.pill{{font-size:19px;font-weight:700;fill:#{NAVY}}}
.flowlabel{{font-size:19px;font-weight:700;fill:#{NAVY}}}
.arrow{{font-size:28px;fill:#{BLUE}}}
.boxlabel{{font-size:24px;font-weight:700;fill:#{WHITE}}}
.boxlabelink{{font-size:24px;font-weight:700;fill:#{NAVY}}}
.boxsub{{font-size:18px;fill:#{MUTED}}}
.boxsubw{{font-size:17px;fill:#{PALE}}}
.tag{{font-size:18px;fill:#{MUTED}}}
.cardhead{{font-size:29px;font-weight:700;fill:#{WHITE}}}
.cardtitle{{font-size:31px;font-weight:700;fill:#{NAVY}}}
.cardsub{{font-size:19px;font-weight:600;fill:#{WHITE}}}
.cardbody{{font-size:19px;fill:#{MUTED}}}
.caption{{font-size:19px;fill:#{MUTED}}}
.small{{font-size:16px;fill:#{MUTED}}}
</style>'''


def base_canvas(body: str, bg: str = BG, bars: bool = True) -> str:
    bar = f'<rect width="{W}" height="14" fill="#{BLUE}"/>' if bars else ""
    return f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}">{STYLE}<rect width="{W}" height="{H}" fill="#{bg}"/>{bar}{body}</svg>'


def render_svg(svg: str, name: str) -> Path:
    svg_path = RENDERS / f"{name}.svg"
    png_path = RENDERS / f"{name}.png"
    svg_path.write_text(svg, encoding="utf-8")
    subprocess.run(["rsvg-convert", "-w", str(W), "-h", str(H), str(svg_path), "-o", str(png_path)], check=True)
    return png_path


# ---------------------------------------------------------------- overlay spec

@dataclass
class Overlay:
    text: str
    x: int
    y: int
    w: int
    h: int
    size: int
    color: str
    bold: bool = False
    align: str = "left"  # left | center | right
    anchor: str = "top"  # top | middle
    line_spacing: float = 1.05


@dataclass
class Slide:
    name: str
    svg_body: str
    overlays: list[Overlay]
    notes: str
    seconds: int
    bg: str = BG


# ---------------------------------------------------------------- slide content

def kicker_overlay() -> Overlay:
    return Overlay("CRYOSTACK", 105, 46, 400, 30, 15, BLUE, bold=True)


def footer_overlay() -> Overlay:
    return Overlay("ICCL + PGSL  ·  GEORGIA TECH", 1300, 1034, 515, 26, 12, MUTED, align="right")


def build_slide1() -> Slide:
    icon = ASSETS_ICON
    icon_w, icon_h = 260, 260 * 628 / 1378
    body = ""
    body += f'<rect width="{W}" height="14" y="{H-14}" fill="#{BLUE}"/>'
    body += image_tag(icon, int((W - icon_w) / 2), 150, int(icon_w), int(icon_h))
    body += hline(int(W / 2 - 210), int(W / 2 + 210), 740, stroke=CYAN, w=3)
    overlays = [
        Overlay("CryoStack", 210, 340, 1500, 120, 88, NAVY, bold=True, align="center"),
        Overlay(
            "A Unified Cyberinfrastructure Platform for Reproducible Ice-Sheet Modeling, "
            "Data Assimilation, and Scientific Computing across HPC and Cloud Environments",
            310, 485, 1300, 110, 22, MUTED, align="center", line_spacing=1.25,
        ),
        Overlay("Brian Kyanjo", 460, 780, 1000, 45, 27, INK, bold=True, align="center"),
        Overlay("Research Scientist II", 460, 828, 1000, 34, 20, MUTED, align="center"),
        Overlay("Georgia Tech · Earth and Atmospheric Sciences", 360, 866, 1200, 34, 20, MUTED, align="center"),
    ]
    notes = (
        "Good morning. I'm Brian Kyanjo, a research scientist in Earth and Atmospheric Sciences at "
        "Georgia Tech. Today I want to introduce CryoStack, a cyberinfrastructure platform for running "
        "ice-sheet modeling, data assimilation, and other cryosphere computing reproducibly across HPC "
        "and cloud systems. In the next ten minutes: why it exists, what it connects, what's working "
        "today, and where you can start."
    )
    return Slide("01-title", base_canvas(body, bars=False), overlays, notes, 30)


def build_slide2() -> Slide:
    body = ""
    inputs = ["Models", "Observations", "Data\nAssimilation", "HPC", "Cloud"]
    n = len(inputs)
    total_w = 1500
    x0 = int((W - total_w) / 2)
    pw = int((total_w - (n - 1) * 24) / n)
    apex_x = W / 2
    apex_y = 470
    for i, label in enumerate(inputs):
        x = x0 + i * (pw + 24)
        cx = x + pw / 2
        body += pill(x, 250, pw, 90, label.replace("\n", " "), fill=PALE, stroke=CYAN, text_cls="pill", r=45)
        body += f'<line x1="{cx}" y1="340" x2="{apex_x}" y2="{apex_y-28}" stroke="#{BORDER}" stroke-width="2"/>'
    body += f'<circle cx="{apex_x}" cy="{apex_y}" r="30" fill="#{NAVY}"/>'
    body += f'<text x="{apex_x}" y="{apex_y+8}" text-anchor="middle" class="cardsub">+</text>'
    body += svg_lines(["one integrated computational demand"], apex_x, apex_y + 62, "caption", 24, "middle")

    steps = ["Install", "Configure", "Port", "Submit", "Monitor", "Retrieve", "Reproduce"]
    body += chevron_flow(steps, 105, 640, 1710, h=100)

    body += hline(105, 1815, 610, stroke=BORDER, w=1)
    return Slide(
        "02-challenge",
        base_canvas(body),
        [
            kicker_overlay(), footer_overlay(),
            Overlay("The challenge is bigger than the model", 105, 90, 1710, 110, 46, NAVY, bold=True),
            Overlay(
                "Every one of those steps is a place where reproducibility can quietly break.",
                105, 790, 1710, 50, 22, MUTED, align="center",
            ),
        ],
        (
            "Cryosphere science increasingly means combining models, observations, and data assimilation "
            "-- and running all of it across HPC and cloud systems. But every one of those pieces carries "
            "its own infrastructure burden: install the software, configure it for a specific machine, "
            "port workflows when you move systems, submit and monitor jobs, retrieve results, and -- if "
            "you're careful -- still be able to reproduce what you did. That burden falls on scientists, "
            "and it grows with how sophisticated our models and observations get."
        ),
        60,
    )


def build_slide3() -> Slide:
    body = ""
    body += pill(410, 380, 1100, 92, "Configure · Launch · Monitor · Visualize · Reproduce",
                 fill=NAVY, stroke=NAVY, text_cls="cardsub", r=20, wrap_chars=60)

    # abstract "environment" diagram: browser chrome + four app tiles + arrow to compute glyph
    bx, by, bw, bh = 260, 560, 1000, 400
    body += f'<rect x="{bx}" y="{by}" width="{bw}" height="{bh}" rx="18" fill="#{WHITE}" stroke="#{BORDER}" stroke-width="2"/>'
    body += f'<rect x="{bx}" y="{by}" width="{bw}" height="52" rx="18" fill="#{PALE}"/>'
    body += f'<rect x="{bx}" y="{by+34}" width="{bw}" height="18" fill="#{PALE}"/>'
    for i, c in enumerate((GREEN, "D99516", "D64242")):
        body += f'<circle cx="{bx+26+i*24}" cy="{by+26}" r="7" fill="#{c}"/>'
    body += f'<rect x="{bx+120}" y="{by+14}" width="{bw-160}" height="26" rx="13" fill="#{WHITE}" stroke="#{BORDER}"/>'
    body += svg_lines(["cryostack.eas.gatech.edu"], bx+140, by+33, "small", 0)
    apps = ["CryoLauncher", "ICESEE", "Frozen\nLegacies", "LIVIST"]
    tile_w = (bw - 40 - 3*20) / 4
    for i, a in enumerate(apps):
        tx = bx + 20 + i * (tile_w + 20)
        ty = by + 90
        body += box(int(tx), int(ty), int(tile_w), 130, a.replace("\n", " "), fill=PALE, stroke=CYAN, label_cls="boxlabelink", r=12)
    body += f'<rect x="{bx+20}" y="{by+250}" width="{bw-40}" height="110" rx="12" fill="#{BG}" stroke="#{BORDER}"/>'
    body += svg_lines(["Shared identity  ·  compute access  ·  job status  ·  results"], bx+bw/2, by+312, "tag", 0, "middle")

    body += f'<text x="{bx+bw+70}" y="{by+bh/2+18}" text-anchor="middle" class="arrow" font-size="46">&#8594;</text>'
    hx, hy = bx+bw+170, by+bh/2-90
    body += box(int(hx), int(hy), 300, 180, "HPC / Cloud", "Your compute", fill=NAVY, stroke=NAVY, label_cls="boxlabel", sub_cls="boxsubw")

    return Slide(
        "03-introducing",
        base_canvas(body),
        [
            kicker_overlay(), footer_overlay(),
            Overlay("Introducing CryoStack", 105, 90, 1710, 110, 46, NAVY, bold=True),
            Overlay("Science first. Infrastructure underneath.", 105, 230, 1710, 70, 34, NAVY, bold=True, align="center"),
        ],
        (
            "This is CryoStack: a single, web-based environment where a researcher configures a "
            "scientific run, launches it, monitors it, visualizes results, and can reproduce the "
            "experiment later -- without becoming a systems administrator along the way. Different "
            "scientific applications connect to that infrastructure in their own way -- I'll show you "
            "that next -- but from the researcher's side it's one consistent environment."
        ),
        60,
    )


def build_slide4() -> Slide:
    body = ""
    cards = [
        ("CryoLauncher", "ISSM · Icepack", "Model configuration and execution", "mountain"),
        ("ICESEE", "Ensemble data assimilation", "State and parameter estimation", "ensemble"),
        ("Frozen Legacies", "Historical Antarctic radar", "Cataloging · interpretation · tracing", "radar"),
        ("LIVIST", "Englacial temperature", "Radar · borehole observations", "thermometer"),
    ]
    n = len(cards)
    total_w = 1710
    cw = int((total_w - (n - 1) * 30) / n)
    y, ch = 300, 560
    for i, (title, sub, body_text, icon_kind) in enumerate(cards):
        x = 105 + i * (cw + 30)
        accent = ACCENTS[i % 4]
        cx = x + cw / 2
        body += f'<rect x="{x}" y="{y}" width="{cw}" height="{ch}" rx="22" fill="#{WHITE}" stroke="#{BORDER}" stroke-width="2"/>'
        body += f'<rect x="{x}" y="{y}" width="{cw}" height="8" rx="4" fill="#{accent}"/>'
        body += app_icon(icon_kind, cx, y + 100, 90, accent)
        body += svg_lines(textwrap.wrap(title, width=16), cx, y+185, "cardtitle", 34, "middle")
        body += hline(int(x+cw/2-40), int(x+cw/2+40), y+210, stroke=BORDER, w=1)
        body += svg_lines(textwrap.wrap(sub, width=17), cx, y+264, "boxlabelink", 28, "middle")
        body += svg_lines(textwrap.wrap(body_text, width=19), cx, y+364, "cardbody", 28, "middle")
    return Slide(
        "04-applications",
        base_canvas(body),
        [
            kicker_overlay(), footer_overlay(),
            Overlay("One platform. Different scientific workflows.", 105, 90, 1710, 110, 44, NAVY, bold=True),
            Overlay("Shared cyberinfrastructure. Application-specific science.", 105, 900, 1710, 60, 26, NAVY, bold=True, align="center"),
        ],
        (
            "CryoStack today hosts four first-class scientific applications. CryoLauncher runs ice-sheet "
            "models -- ISSM and Icepack. ICESEE performs ensemble data assimilation for state and "
            "parameter estimation. Frozen Legacies works with historical Antarctic radar -- cataloging, "
            "interpreting, and tracing. LIVIST focuses on englacial temperature from radar and borehole "
            "observations. These are genuinely different workflows -- CryoStack doesn't force one "
            "universal pipeline on them. What they share is the cyberinfrastructure underneath: identity, "
            "compute access, and a common web environment."
        ),
        90,
    )


def build_slide5() -> Slide:
    body = ""
    cx = W / 2
    top_y, top_h = 200, 90
    body += box(int(cx-160), top_y, 320, top_h, "CryoStack", fill=NAVY, stroke=NAVY, label_cls="boxlabel")

    rail_y = top_y + top_h + 50
    branch_y, branch_h = rail_y + 40, 120
    xs = {"Local": 330, "Remote": cx, "Cloud": 1590}
    body += vline(int(cx), top_y+top_h, rail_y, stroke=BORDER, w=2)
    body += hline(xs["Local"], xs["Cloud"], rail_y, stroke=BORDER, w=2)
    for x in xs.values():
        body += vline(int(x), rail_y, branch_y, stroke=BORDER, w=2)

    body += box(int(xs["Local"]-160), branch_y, 320, branch_h, "Local", "where supported (e.g., ICESEE)", fill=WHITE, stroke=CYAN, label_cls="boxlabelink", sub_wrap=22)
    body += box(int(xs["Remote"]-160), branch_y, 320, branch_h, "Remote", fill=WHITE, stroke=CYAN, label_cls="boxlabelink")
    body += box(int(xs["Cloud"]-160), branch_y, 320, branch_h, "Cloud", fill=WHITE, stroke=CYAN, label_cls="boxlabelink")

    row2_y, row2_h = branch_y + branch_h + 60, 150
    body += vline(int(xs["Remote"]), branch_y+branch_h, row2_y, stroke=BORDER, w=2)
    body += box(int(xs["Remote"]-210), row2_y, 420, row2_h, "Direct SSH  /  Connector-Relay", sub="connectivity, not an execution backend", fill=PALE, stroke=CYAN, label_cls="boxlabelink", sub_wrap=24)

    body += vline(int(xs["Cloud"]), branch_y+branch_h, row2_y, stroke=BORDER, w=2)
    body += box(int(xs["Cloud"]-190), row2_y, 380, row2_h, "User-managed AWS account", fill=PALE, stroke=CYAN, label_cls="boxlabelink")

    row3_y, row3_h = row2_y + row2_h + 60, 140
    body += vline(int(xs["Remote"]), row2_y+row2_h, row3_y, stroke=BORDER, w=2)
    body += box(int(xs["Remote"]-210), row3_y, 420, row3_h, "Institutional / National HPC", "PACE · other configured HPC systems", fill=WHITE, stroke=BORDER, label_cls="boxlabelink", sub_wrap=24)

    body += vline(int(xs["Cloud"]), row2_y+row2_h, row3_y, stroke=BORDER, w=2)
    body += box(int(xs["Cloud"]-190), row3_y, 380, row3_h, "Fargate / EC2 On-Demand", fill=WHITE, stroke=BORDER, label_cls="boxlabelink")

    divider_y = row3_y + row3_h + 40
    body += hline(105, 1815, divider_y, stroke=BORDER, w=1)
    return Slide(
        "05-architecture",
        base_canvas(body),
        [
            kicker_overlay(), footer_overlay(),
            Overlay("From your browser to the compute you already use", 105, 90, 1710, 110, 40, NAVY, bold=True),
            Overlay(
                "CryoLauncher supports Remote and Cloud. Local applies where supported, particularly ICESEE.",
                105, divider_y+20, 1710, 50, 19, MUTED, align="center",
            ),
        ],
        (
            "CryoStack connects to compute you already have access to. For applications where it makes "
            "sense -- particularly ICESEE -- you can run locally. For remote HPC, CryoStack connects "
            "either through direct SSH or through our Connector/Relay -- connectivity infrastructure, not "
            "a separate execution backend -- either way landing on institutional or national HPC like "
            "PACE. For cloud, CryoStack currently uses AWS, on Fargate or EC2 on-demand, under your own "
            "account. CryoLauncher specifically supports Remote and Cloud, not Local. The point: CryoStack "
            "doesn't replace your compute, it gives you one consistent way to reach it."
        ),
        90,
    )


def build_slide6() -> Slide:
    body = ""
    col_w = 800
    lx, rx = 105, 105 + col_w + 110
    y, img_h = 220, 560
    body += image_tag(ASSETS / "issm-stressbalance-velocity.png", lx, y, col_w, img_h)
    body += f'<rect x="{lx}" y="{y}" width="{col_w}" height="{img_h}" rx="10" fill="none" stroke="#{BORDER}" stroke-width="2"/>'
    body += svg_lines(["ISSM -- Remote execution (PACE via Connector)"], lx, y+img_h+42, "cardtitle", 0)

    body += image_tag(ASSETS / "icesee-lorenz96-result.png", rx, y, col_w, img_h)
    body += f'<rect x="{rx}" y="{y}" width="{col_w}" height="{img_h}" rx="10" fill="none" stroke="#{BORDER}" stroke-width="2"/>'
    body += svg_lines(["ICESEE -- ensemble data assimilation"], rx, y+img_h+42, "cardtitle", 0)

    pill_y = y + img_h + 80
    body += pill(int(W/2-340), pill_y, 680, 76, "Also demonstrated on AWS: ISSM & Icepack -- Fargate / EC2 On-Demand", fill=PALE, stroke=CYAN, text_cls="pill", r=16, wrap_chars=34)
    return Slide(
        "06-real-science",
        base_canvas(body),
        [
            kicker_overlay(), footer_overlay(),
            Overlay("It is already running real science", 105, 90, 1710, 110, 44, NAVY, bold=True),
        ],
        (
            "This isn't a proposed architecture -- it's running today. CryoLauncher has executed ISSM on "
            "institutional HPC and on AWS Batch, both Fargate and EC2 on-demand, including full MPI solver "
            "runs, postprocessing, retrieval, and visualization. Icepack has been demonstrated on those same "
            "cloud paths. ICESEE has its own local, remote, and cloud execution paths, and we've demonstrated "
            "its Lorenz-96 configuration running in the cloud. When a run needs Georgia Tech's institutional "
            "MATLAB license, CryoLauncher can reach it from AWS through Connector/Relay. To be precise about "
            "where we are: Spot instances, GPUs, multi-node execution, and a fully cloud-coupled ICESEE-ISSM "
            "workflow are not yet validated -- that's ongoing work, not a claim I'm making today."
        ),
        90,
    )


def build_slide7() -> Slide:
    body = ""
    stages = ["Scientific\nconfiguration", "Computational\nenvironment", "Execution +\nmetadata", "Results +\nprovenance", "Reproduce ·\nCompare · Share"]
    n = len(stages)
    total_w = 1710
    x0 = 105
    bw = int((total_w - (n-1)*40) / n)
    y = 440
    for i, s in enumerate(stages):
        x = x0 + i*(bw+40)
        fill = NAVY if i == n-1 else WHITE
        lbl_cls = "boxlabel" if i == n-1 else "boxlabelink"
        stroke = NAVY if i == n-1 else CYAN
        body += box(x, y, bw, 170, s.replace("\n", " "), fill=fill, stroke=stroke, label_cls=lbl_cls, r=16)
        if i < n-1:
            ax = x+bw+20
            body += f'<text x="{ax}" y="{y+95}" text-anchor="middle" class="arrow" font-size="34">&#8594;</text>'
    body += hline(105, 1815, y+230, stroke=BORDER, w=1)
    return Slide(
        "07-reproducibility",
        base_canvas(body),
        [
            kicker_overlay(), footer_overlay(),
            Overlay("The experiment is more than its final output", 105, 90, 1710, 110, 42, NAVY, bold=True),
            Overlay(
                "CryoStack is built to carry an experiment's context through its lifecycle -- not just the final result.",
                105, 720, 1710, 60, 22, MUTED, align="center",
            ),
        ],
        (
            "A result is only as useful as your ability to trust and repeat it. That's the goal behind "
            "CryoStack's reproducibility architecture: carry an experiment's scientific configuration, "
            "computational environment, execution metadata, and results with provenance across its "
            "lifecycle, so it can be reproduced, compared, and shared. How far that goes today varies by "
            "application -- it's further along for some workflows than others -- but it's core to why "
            "CryoStack exists, not an afterthought."
        ),
        60,
    )


def build_slide8() -> Slide:
    body = ""
    body += f'<rect width="{W}" height="14" y="{H-14}" fill="#{BLUE}"/>'
    icon = ASSETS_ICON
    icon_w = 110
    icon_h = icon_w * 628 / 1378
    body += image_tag(icon, 105, 95, int(icon_w), int(icon_h))

    qr = ASSETS / "qr-cryostack.png"
    qs = 380
    qx, qy = 1815 - qs - 80, 430
    body += f'<rect x="{qx-26}" y="{qy-26}" width="{qs+52}" height="{qs+52}" rx="22" fill="#{WHITE}" stroke="#{CYAN}" stroke-width="3"/>'
    body += image_tag(qr, qx, qy, qs, qs)
    body += svg_lines(["SCAN TO EXPLORE"], qx+qs/2, qy-42, "pill", 0, "middle")

    body += svg_lines(
        ["More models · datasets · analysis tools · community workflows  --  next: PISM, Elmer/Ice, SICOPOLIS"],
        W/2, 990, "small", 0, "middle")
    return Slide(
        "08-live",
        base_canvas(body),
        [
            Overlay("CryoStack is live.", 105, 240, 1500, 130, 80, NAVY, bold=True),
            Overlay("Explore the platform. Run a workflow. Help us build what comes next.", 105, 380, 1500, 70, 26, MUTED),
            Overlay("cryostack.eas.gatech.edu", 105, 500, 1000, 70, 42, BLUE, bold=True),
            Overlay(f"Documentation  ·  {DOCS_URL}", 105, 580, 1200, 40, 20, MUTED),
        ],
        (
            "We're actively growing this: more models like PISM, Elmer/Ice, and SICOPOLIS are on our "
            "onboarding roadmap, along with more datasets, analysis tools, and community workflows. But "
            "the main point today is simpler than that: CryoStack is live, right now, at this address -- "
            "and we invite you to explore it, run a workflow, and help us shape what comes next. Thank you."
        ),
        60,
    )


ASSETS_ICON = ASSETS / "cryostack-icon-only.png"


def build_slide9() -> Slide:
    body = ""
    n = len(DEMO_PANELS)
    total_w = 1710
    gap = 40
    cw = int((total_w - (n - 1) * gap) / n)
    thumb_pad = 20
    thumb_w = cw - 2 * thumb_pad
    thumb_h = int(thumb_w * 9 / 16)
    card_h = thumb_h + 2 * thumb_pad + 90
    region_top, region_bottom = 280, 980
    block_h = card_h + 55 + 50  # cards + gap-to-caption + caption line
    card_y = region_top + int((region_bottom - region_top - block_h) / 2)

    for i, (folder, label) in enumerate(DEMO_PANELS):
        x = 105 + i * (cw + gap)
        body += f'<rect x="{x}" y="{card_y}" width="{cw}" height="{card_h}" rx="20" fill="#{WHITE}" stroke="#{BORDER}" stroke-width="2"/>'
        tx, ty = x + thumb_pad, card_y + thumb_pad
        body += image_tag(demo_thumbnail_path(folder), tx, ty, thumb_w, thumb_h)
        body += f'<rect x="{tx}" y="{ty}" width="{thumb_w}" height="{thumb_h}" rx="10" fill="none" stroke="#{BORDER}" stroke-width="2"/>'
        body += play_button(tx + thumb_w / 2, ty + thumb_h / 2, thumb_h * 0.16)
        body += svg_lines(textwrap.wrap(label, width=20), x + cw / 2, ty + thumb_h + 50, "cardtitle", 32, "middle")

    caption_y = card_y + card_h + 55
    body += hline(105, 1815, caption_y - 30, stroke=BORDER, w=1)
    return Slide(
        "09-demo",
        base_canvas(body),
        [
            kicker_overlay(), footer_overlay(),
            Overlay("See CryoStack in action", 105, 90, 1710, 90, 44, NAVY, bold=True),
            Overlay("From configuration to scientific results", 105, 195, 1710, 50, 24, MUTED),
            Overlay(
                "Short excerpts here -- full recordings will be published as CryoStack tutorials.",
                105, caption_y, 1710, 50, 19, MUTED, align="center",
            ),
        ],
        (
            "Rather than only showing the architecture, I want to briefly show CryoStack actually "
            "running scientific workflows. Here are three examples: ISSM running remotely on HPC, ISSM "
            "running in the cloud, and Icepack running in the cloud. In each case, the researcher works "
            "through CryoLauncher while CryoStack handles the underlying execution. These are short "
            "excerpts -- the full demonstrations will also be published as CryoStack tutorials."
        ),
        50,
    )


def all_slides() -> list[Slide]:
    return [
        build_slide1(), build_slide2(), build_slide3(), build_slide4(),
        build_slide5(), build_slide6(), build_slide7(), build_slide8(),
        build_slide9(),
    ]


# ---------------------------------------------------------------- text fitting

def fit_size(text: str, max_w_px: float, max_h_px: float, start_size: int, bold: bool, min_size: int = 14, line_spacing: float = 1.05):
    font_path = FONT_BOLD if bold else FONT_REG
    min_size = min(min_size, start_size)
    size = start_size
    while size >= min_size:
        font = ImageFont.truetype(font_path, size)
        words = text.split()
        lines, cur = [], ""
        for word in words:
            trial = (cur + " " + word).strip()
            if font.getlength(trial) <= max_w_px:
                cur = trial
            else:
                if cur:
                    lines.append(cur)
                cur = word
        if cur:
            lines.append(cur)
        total_h = len(lines) * size * line_spacing
        if total_h <= max_h_px:
            return size, lines, font
        size -= 2
    return size, lines, ImageFont.truetype(font_path, size)


# ---------------------------------------------------------------- PPTX build

def build_pptx(slides: list[Slide], bg_pngs: list[Path], out_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Emu(W * PX_TO_EMU)
    prs.slide_height = Emu(H * PX_TO_EMU)
    blank = prs.slide_layouts[6]

    for slide_spec, png in zip(slides, bg_pngs):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(png), Emu(0), Emu(0), width=Emu(W * PX_TO_EMU), height=Emu(H * PX_TO_EMU))
        for ov in slide_spec.overlays:
            size, lines, _ = fit_size(ov.text, ov.w, ov.h, ov.size, ov.bold, line_spacing=ov.line_spacing)
            tb = slide.shapes.add_textbox(Emu(ov.x * PX_TO_EMU), Emu(ov.y * PX_TO_EMU), Emu(ov.w * PX_TO_EMU), Emu(ov.h * PX_TO_EMU))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.TOP if ov.anchor == "top" else MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[ov.align]
            r = p.add_run()
            r.text = ov.text
            r.font.size = Pt(size)
            r.font.bold = ov.bold
            r.font.name = "Arial"
            r.font.color.rgb = hexrgb(ov.color)
        notes = slide.notes_slide
        notes.notes_text_frame.text = slide_spec.notes
    prs.save(str(out_path))


# ---------------------------------------------------------------- QA compositor + PDF

def composite_preview(slide_spec: Slide, bg_png: Path) -> Image.Image:
    im = Image.open(bg_png).convert("RGB")
    draw = ImageDraw.Draw(im)
    for ov in slide_spec.overlays:
        size, lines, font = fit_size(ov.text, ov.w, ov.h, ov.size, ov.bold, line_spacing=ov.line_spacing)
        color = tuple(int(ov.color[i:i+2], 16) for i in (0, 2, 4))
        line_h = size * ov.line_spacing
        total_h = len(lines) * line_h
        y = ov.y if ov.anchor == "top" else ov.y + (ov.h - total_h) / 2
        for line in lines:
            w = font.getlength(line)
            if ov.align == "center":
                x = ov.x + (ov.w - w) / 2
            elif ov.align == "right":
                x = ov.x + ov.w - w
            else:
                x = ov.x
            draw.text((x, y), line, font=font, fill=color)
            y += line_h
    return im


def build_previews_and_pdf(slides: list[Slide], bg_pngs: list[Path]) -> None:
    images = []
    for slide_spec, png in zip(slides, bg_pngs):
        im = composite_preview(slide_spec, png)
        im.save(PREVIEWS / f"{slide_spec.name}.png")
        images.append(im)
    pdf_path = OUT / "output" / "CryoStack_Launch_Deck.pdf"
    images[0].save(pdf_path, save_all=True, append_images=images[1:])


def main() -> None:
    prepare_assets()
    prepare_media_placeholders()
    slides = all_slides()
    bg_pngs = [render_svg(s.svg_body, s.name) for s in slides]
    build_pptx(slides, bg_pngs, OUT / "output" / "CryoStack_Launch_Deck.pptx")
    build_previews_and_pdf(slides, bg_pngs)
    total = sum(s.seconds for s in slides)
    print(f"Built {len(slides)} slides. Estimated total speaking time: {total}s (~{total/60:.1f} min)")
    for s in slides:
        print(f"  {s.name}: ~{s.seconds}s")


if __name__ == "__main__":
    main()
